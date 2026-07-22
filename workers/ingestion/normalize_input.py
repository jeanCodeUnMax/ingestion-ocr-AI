#!/usr/bin/env python3
"""Normalize supported inputs into a canonical PDF and Markdown sidecar.
The canonical PDF feeds the existing page inventory/OCR pipeline; the original file is preserved separately.
"""
from __future__ import annotations

import argparse
import csv
import json
import mimetypes
import os
import shutil
import sys
import textwrap
from pathlib import Path
from typing import Iterable

import fitz


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def markdown_table(rows: list[list[object]]) -> str:
    if not rows:
        return "[Feuille vide]"
    width = max(len(row) for row in rows)
    normalized = [[clean_cell(row[i]) if i < len(row) else "" for i in range(width)] for row in rows]
    header = normalized[0]
    body = normalized[1:]
    return "\n".join([
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
        *["| " + " | ".join(row) + " |" for row in body],
    ])


def extract_docx(source: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx est requis pour les fichiers Word.") from exc
    document = Document(source)
    blocks: list[str] = [f"# {source.stem}", ""]
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            try:
                level = max(1, min(6, int(style.split()[-1])))
            except Exception:
                level = 2
            blocks.extend([f"{'#' * level} {text}", ""])
        elif style.startswith("list"):
            blocks.append(f"- {text}")
        else:
            blocks.extend([text, ""])
    for index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        blocks.extend([f"## Tableau {index}", "", markdown_table(rows), ""])
    return "\n".join(blocks).strip() + "\n"


def extract_xlsx(source: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl est requis pour les fichiers Excel.") from exc
    workbook = load_workbook(source, data_only=False, read_only=True)
    blocks: list[str] = [f"# Classeur {source.stem}", ""]
    for sheet in workbook.worksheets:
        blocks.extend([f"## Feuille : {sheet.title}", ""])
        rows: list[list[object]] = []
        for row in sheet.iter_rows(values_only=True):
            values = list(row)
            if any(value not in (None, "") for value in values):
                rows.append(values)
        blocks.extend([markdown_table(rows), ""])
    return "\n".join(blocks).strip() + "\n"


def extract_csv(source: Path) -> str:
    raw = source.read_text(encoding="utf-8-sig", errors="replace")
    sample = raw[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(raw.splitlines(), dialect=dialect))
    return f"# {source.stem}\n\n{markdown_table(rows)}\n"


def extract_json(source: Path) -> str:
    value = json.loads(source.read_text(encoding="utf-8-sig"))
    return f"# {source.stem}\n\n```json\n{json.dumps(value, indent=2, ensure_ascii=False)}\n```\n"


def extract_text(source: Path, kind: str) -> str:
    text = source.read_text(encoding="utf-8-sig", errors="replace")
    if kind == "markdown":
        return text
    return f"# {source.stem}\n\n{text}\n"


def image_to_pdf(source: Path, output_pdf: Path) -> None:
    image_document = fitz.open(source)
    pdf_bytes = image_document.convert_to_pdf()
    image_document.close()
    pdf = fitz.open("pdf", pdf_bytes)
    pdf.save(output_pdf)
    pdf.close()


def render_markdown_to_pdf(markdown: str, output_pdf: Path) -> int:
    # A deterministic, dependency-light renderer. The Markdown source remains the canonical text sidecar.
    page_width, page_height = fitz.paper_size("a4")
    margin = 48
    font_size = 10.5
    line_height = 14
    usable_chars = 105
    lines: list[str] = []
    for raw_line in markdown.splitlines():
        if not raw_line:
            lines.append("")
            continue
        prefix = ""
        content = raw_line
        if raw_line.startswith("#"):
            hashes = len(raw_line) - len(raw_line.lstrip("#"))
            content = raw_line[hashes:].strip().upper() if hashes <= 2 else raw_line[hashes:].strip()
            prefix = "" if hashes <= 2 else "  "
        elif raw_line.startswith("- "):
            prefix, content = "• ", raw_line[2:]
        elif raw_line.startswith("| "):
            prefix = ""
        wrapped = textwrap.wrap(prefix + content, width=usable_chars, replace_whitespace=False, drop_whitespace=False) or [""]
        lines.extend(wrapped)

    doc = fitz.open()
    max_lines = max(1, int((page_height - 2 * margin) // line_height))
    page_count = 0
    for start in range(0, len(lines) or 1, max_lines):
        page = doc.new_page(width=page_width, height=page_height)
        page_count += 1
        segment = lines[start:start + max_lines]
        y = margin
        for line in segment:
            page.insert_text((margin, y), line, fontsize=font_size, fontname="helv")
            y += line_height
    if page_count == 0:
        doc.new_page(width=page_width, height=page_height)
        page_count = 1
    doc.set_metadata({"title": "OCR AI System canonical document", "producer": "OCR AI System v0.5"})
    doc.save(output_pdf)
    doc.close()
    return page_count


def detect_kind(source: Path, requested: str | None) -> str:
    if requested:
        return requested
    ext = source.suffix.lower()
    mapping = {
        ".pdf": "pdf", ".docx": "word", ".xlsx": "excel", ".xlsm": "excel",
        ".csv": "csv", ".json": "json", ".md": "markdown", ".markdown": "markdown",
        ".txt": "text", ".png": "image", ".jpg": "image", ".jpeg": "image",
        ".webp": "image", ".tif": "image", ".tiff": "image", ".svg": "svg",
        ".wav": "audio", ".mp3": "audio", ".m4a": "audio", ".ogg": "audio",
    }
    return mapping.get(ext, "text")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--kind")
    args = parser.parse_args()

    source = Path(args.input).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    kind = detect_kind(source, args.kind)
    canonical_pdf = output_dir / "canonical.pdf"
    canonical_md = output_dir / "canonical_source.md"

    if kind == "audio":
        raise RuntimeError("AUDIO_TRANSCRIPTION_PROVIDER_REQUIRED")
    if kind == "pdf":
        shutil.copy2(source, canonical_pdf)
        canonical_md.write_text("", encoding="utf-8")
        pages = fitz.open(canonical_pdf).page_count
    elif kind in {"image", "svg"}:
        image_to_pdf(source, canonical_pdf)
        canonical_md.write_text(
            f"# Image {source.name}\n\nL'image doit être décrite par le moteur OCR/vision.\n",
            encoding="utf-8",
        )
        pages = fitz.open(canonical_pdf).page_count
    else:
        if kind == "word":
            markdown = extract_docx(source)
        elif kind == "excel":
            markdown = extract_xlsx(source)
        elif kind == "csv":
            markdown = extract_csv(source)
        elif kind == "json":
            markdown = extract_json(source)
        elif kind in {"markdown", "text"}:
            markdown = extract_text(source, kind)
        else:
            raise RuntimeError(f"INPUT_KIND_UNSUPPORTED:{kind}")
        canonical_md.write_text(markdown, encoding="utf-8")
        pages = render_markdown_to_pdf(markdown, canonical_pdf)

    payload = {
        "schemaVersion": "1.0",
        "kind": kind,
        "sourcePath": str(source),
        "canonicalPdfPath": str(canonical_pdf),
        "canonicalMarkdownPath": str(canonical_md),
        "pageCount": pages,
        "mimeType": mimetypes.guess_type(source.name)[0] or "application/octet-stream",
    }
    print(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        raise
